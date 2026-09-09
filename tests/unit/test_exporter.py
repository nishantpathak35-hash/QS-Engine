"""
Unit Tests for Multi-Format Exporter (Excel, CSV, JSON)
"""

import pytest
from pathlib import Path
import openpyxl
from core.models.takeoff import TakeoffSummary, TakeoffLineItem
from core.units import DisplayUnit
from core.models.semantics import EntityStatus
from exports.exporter import QSExporter

def test_export_excel_structure(tmp_path: Path):
    takeoff = TakeoffSummary(
        drawing_id="DRW-TEST-01",
        drawing_number="A-101",
        revision="03",
        items=[
            TakeoffLineItem(
                id="TO-001",
                item_code="FL-01",
                description="Vitrified Tile Flooring",
                location="Conference Room",
                quantity=24.0,
                unit=DisplayUnit.SQM,
                formula="6m * 4m = 24.000 sqm",
                confidence=0.99,
                status=EntityStatus.VERIFIED
            ),
            TakeoffLineItem(
                id="TO-002",
                item_code="SK-01",
                description="Skirting",
                location="Conference Room",
                quantity=19.1,
                unit=DisplayUnit.M,
                formula="20m - 0.9m = 19.100 m",
                confidence=0.98,
                status=EntityStatus.AUTO_MEASURED
            )
        ]
    )

    excel_file = tmp_path / "test_takeoff.xlsx"
    QSExporter.export_excel(takeoff, excel_file)

    assert excel_file.exists()
    assert excel_file.stat().st_size > 1000

    # Load with openpyxl to verify sheets & data
    wb = openpyxl.load_workbook(str(excel_file))
    assert "Measurement Takeoff" in wb.sheetnames
    assert "BOQ Summary" in wb.sheetnames

    ws1 = wb["Measurement Takeoff"]
    assert ws1["A1"].value == "QUANTIFICATION ENGINE — DETAILED MEASUREMENT TAKEOFF"
    # Row 5 is first data row
    assert ws1.cell(row=5, column=1).value == "FL-01"
    assert ws1.cell(row=5, column=4).value == 24.0

    ws2 = wb["BOQ Summary"]
    assert ws2.cell(row=5, column=1).value == "FL-01"
    assert ws2.cell(row=5, column=3).value == 24.0

def test_export_csv_and_json():
    takeoff = TakeoffSummary(
        drawing_id="DRW-TEST-02",
        drawing_number="A-102",
        revision="01",
        items=[
            TakeoffLineItem(
                id="TO-001",
                item_code="FL-01",
                description="Flooring",
                location="Office",
                quantity=30.5,
                unit=DisplayUnit.SQM,
                formula="30.500 sqm",
                confidence=0.97
            )
        ]
    )

    csv_str = QSExporter.export_csv(takeoff)
    assert "FL-01" in csv_str
    assert "30.5" in csv_str
    assert "BOQ TOTALS SUMMARY" in csv_str

    json_str = QSExporter.export_json(takeoff)
    assert '"drawing_number": "A-102"' in json_str
    assert '"FL-01"' in json_str
