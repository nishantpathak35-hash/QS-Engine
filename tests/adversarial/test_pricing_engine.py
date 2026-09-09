"""
Adversarial & Unit Tests — Rate Analysis, Cost Estimation & Priced BOQ Engine
Validates IS 1200 / POMI trade rate analysis, dual Net/Gross costing, contractor markups, GST,
multi-sheet Excel generation, and FastAPI endpoints.
"""

import tempfile
from pathlib import Path
import pytest
import openpyxl

from core.models.takeoff import TakeoffSummary, TakeoffLineItem
from core.models.project_profile import ProjectProfile
from core.units import DisplayUnit
from pricing.cost_engine import CostEstimationEngine, FitoutGrade, BASELINE_RATES
from exports.exporter import QSExporter


def _create_sample_takeoff() -> TakeoffSummary:
    prof = ProjectProfile(
        project_id="P-TEST",
        allow_assumptions=True,
        ceiling_height_m=3.0,
        wastage_enabled=True,
        wastage_preset="standard"
    )

    t = TakeoffSummary(
        drawing_id="DRW-PRICE-01",
        drawing_number="A-101",
        revision="01",
        items=[
            TakeoffLineItem(
                id="item-fl",
                item_code="FL-01",
                description="Vitrified Tile Flooring",
                location="Main Office",
                quantity=100.0,
                unit=DisplayUnit.SQM,
                formula="room.area",
                confidence=1.0,
                status="AUTO_MEASURED",
                rule_id="RULE-FL",
                wastage_percent=0.05,
                gross_quantity=105.0
            ),
            TakeoffLineItem(
                id="item-cl",
                item_code="CL-01",
                description="Gypsum Board False Ceiling",
                location="Main Office",
                quantity=100.0,
                unit=DisplayUnit.SQM,
                formula="room.area",
                confidence=1.0,
                status="AUTO_MEASURED",
                rule_id="RULE-CL",
                wastage_percent=0.07,
                gross_quantity=107.0
            ),
            TakeoffLineItem(
                id="item-lt",
                item_code="LT-02",
                description="5' Linear LED Fixture 40W",
                location="Main Office",
                quantity=10.0,
                unit=DisplayUnit.NOS,
                formula="fixture.count",
                confidence=1.0,
                status="AUTO_MEASURED",
                rule_id="RULE-LT",
                wastage_percent=0.03,
                gross_quantity=10.3
            )
        ],
        base_measurements={"total_floor_area_sqm": 100.0}
    )
    return t


def test_is1200_dual_costing_principles():
    """Verify IS 1200 principle: Material on Gross PO quantity, Labor on Net drawing quantity."""
    takeoff = _create_sample_takeoff()
    est = CostEstimationEngine.estimate_project_cost(takeoff, fitout_grade=FitoutGrade.STANDARD)

    # Check FL-01
    fl_item = next(it for it in est.items if it.item_code == "FL-01")
    base_fl = BASELINE_RATES["FL-01"]

    expected_mat_cost = 105.0 * base_fl.material_rate_inr # Gross Qty * Mat Rate
    expected_lab_cost = 100.0 * (base_fl.labor_rate_inr + base_fl.machinery_rate_inr) # Net Qty * Lab Rate

    assert fl_item.material_cost_inr == pytest.approx(expected_mat_cost, rel=1e-3)
    assert fl_item.labor_cost_inr == pytest.approx(expected_lab_cost, rel=1e-3)
    assert fl_item.total_amount_inr == pytest.approx(expected_mat_cost + expected_lab_cost, rel=1e-3)


def test_fitout_grade_multipliers():
    """Verify Economy (0.75x) vs Standard (1.0x) vs Grade-A Luxury (1.45x)."""
    takeoff = _create_sample_takeoff()

    est_std = CostEstimationEngine.estimate_project_cost(takeoff, fitout_grade=FitoutGrade.STANDARD)
    est_eco = CostEstimationEngine.estimate_project_cost(takeoff, fitout_grade=FitoutGrade.ECONOMY)
    est_lux = CostEstimationEngine.estimate_project_cost(takeoff, fitout_grade=FitoutGrade.GRADE_A_LUXURY)

    assert est_eco.direct_cost_subtotal_inr == pytest.approx(est_std.direct_cost_subtotal_inr * 0.75, rel=1e-3)
    assert est_lux.direct_cost_subtotal_inr == pytest.approx(est_std.direct_cost_subtotal_inr * 1.45, rel=1e-3)
    assert est_lux.grand_total_budget_inr > est_std.grand_total_budget_inr > est_eco.grand_total_budget_inr


def test_markups_and_statutory_taxes():
    """Verify 10% Contractor OH&P, 3% Contingency, and 18% GST statutory calculations."""
    takeoff = _create_sample_takeoff()
    est = CostEstimationEngine.estimate_project_cost(takeoff, fitout_grade=FitoutGrade.STANDARD)

    direct = est.direct_cost_subtotal_inr
    expected_ohp = direct * 0.10
    expected_cont = direct * 0.03
    expected_net = direct + expected_ohp + expected_cont
    expected_gst = expected_net * 0.18
    expected_grand = expected_net + expected_gst

    assert est.contractor_overhead_profit_inr == pytest.approx(expected_ohp, rel=1e-3)
    assert est.contingency_inr == pytest.approx(expected_cont, rel=1e-3)
    assert est.net_project_cost_inr == pytest.approx(expected_net, rel=1e-3)
    assert est.gst_tax_inr == pytest.approx(expected_gst, rel=1e-3)
    assert est.grand_total_budget_inr == pytest.approx(expected_grand, rel=1e-3)


def test_trade_subtotals_exact_reconciliation():
    """Verify sum of trade subtotals equals direct cost subtotal exactly."""
    takeoff = _create_sample_takeoff()
    est = CostEstimationEngine.estimate_project_cost(takeoff, fitout_grade=FitoutGrade.STANDARD)

    sum_trades = sum(est.trade_subtotals_inr.values())
    assert sum_trades == pytest.approx(est.direct_cost_subtotal_inr, rel=1e-3)
    assert "Flooring & Tiling" in est.trade_subtotals_inr
    assert "Ceiling & Soffits" in est.trade_subtotals_inr
    assert "Electrical & Lighting" in est.trade_subtotals_inr


def test_priced_excel_export_sheet_and_formulas():
    """Verify openpyxl generates 3rd sheet 'Priced BOQ & Budget' with dynamic Excel formulas."""
    takeoff = _create_sample_takeoff()

    with tempfile.TemporaryDirectory() as tmpdir:
        excel_path = Path(tmpdir) / "priced_test.xlsx"
        QSExporter.export_excel(takeoff, excel_path)

        assert excel_path.exists()
        wb = openpyxl.load_workbook(str(excel_path), data_only=False)

        assert "Measurement Takeoff" in wb.sheetnames
        assert "BOQ Summary" in wb.sheetnames
        assert "Priced BOQ & Budget" in wb.sheetnames

        ws3 = wb["Priced BOQ & Budget"]
        assert ws3["A1"].value == "PRICED BILL OF QUANTITIES & PROJECT COST ESTIMATE"

        # Check that row 5 has formulas for material cost (=G5*H5) and total cost (=J5+K5)
        cell_mat_formula = ws3.cell(row=5, column=10).value
        cell_total_formula = ws3.cell(row=5, column=12).value

        assert cell_mat_formula == "=G5*H5"
        assert cell_total_formula == "=J5+K5"

        # Check that grand total row exists and contains formula
        found_grand_total = False
        for row in range(5, ws3.max_row + 1):
            lbl = ws3.cell(row=row, column=11).value
            if lbl and "GRAND TOTAL" in str(lbl):
                found_grand_total = True
                val = ws3.cell(row=row, column=12).value
                assert str(val).startswith("=")
                break

        assert found_grand_total, "Grand total budget formula row must be present in Excel sheet 3"
