import sys
sys.path.insert(0, '.')
from vision.ai_drawing_agent import AIDrawingAgent
from qs.rule_engine import QSRuleEngine
from exports.exporter import QSExporter
import openpyxl

path = 'C:/Users/Admin/Downloads/DENTON LINK LEGAL FINAL 01-FURNITURE LAYOUT.pdf'
insight = AIDrawingAgent.analyze_pdf(path)
rooms = AIDrawingAgent.generate_semantic_rooms(insight, 'DWG-DENTON', insight.scale_mm_per_pt)
ai_items = AIDrawingAgent.generate_takeoff_items(insight, 'DWG-DENTON', 'DENTON-01')

rule_engine = QSRuleEngine()
takeoff = rule_engine.calculate_takeoff('DWG-DENTON', 'DENTON-01', '0', rooms, [], [])
takeoff.items = [it for it in takeoff.items if it.item_code not in ('FL-RAW', 'CL-RAW', 'PT-RAW', 'PT-01')]
takeoff.items.extend(ai_items)

out_xlsx = 'scratch/test_denton_export.xlsx'
QSExporter.export_excel(takeoff, out_xlsx)

wb = openpyxl.load_workbook(out_xlsx, data_only=False)
print('Sheet names:', wb.sheetnames)
ws1 = wb['Measurement Takeoff']
print('Takeoff rows:', ws1.max_row)
ws2 = wb['BOQ Summary']
print('Summary rows:', ws2.max_row)
ws3 = wb['Priced BOQ & Budget']
print('Priced rows:', ws3.max_row)
for r in range(ws3.max_row - 7, ws3.max_row + 1):
    print(f"Row {r}: K={ws3.cell(row=r, column=11).value} | L={ws3.cell(row=r, column=12).value}")
