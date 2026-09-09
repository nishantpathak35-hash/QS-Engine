"""
QS Quantification Engine — Standardized Multi-Format Exporter
Generates auditable Excel (.xlsx), CSV, and JSON measurement sheets.
Enforces Blueprint Section 44 (PDF/Measurement Sheet) & Section 89 (Output Formats).
"""

from __future__ import annotations
import json
import csv
from io import StringIO
from pathlib import Path
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

from core.models.takeoff import TakeoffSummary

class QSExporter:
    """Serializes TakeoffSummary into industry-standard QS formats."""

    @staticmethod
    def export_json(takeoff: TakeoffSummary, indent: int = 2) -> str:
        """Serializes takeoff to structured JSON payload."""
        payload = {
            "drawing_id": takeoff.drawing_id,
            "drawing_number": takeoff.drawing_number,
            "revision": takeoff.revision,
            "items": [item.to_dict() for item in takeoff.items],
            "totals": takeoff.total_by_item(),
            "exceptions": takeoff.exceptions
        }
        return json.dumps(payload, indent=indent)

    @staticmethod
    def export_csv(takeoff: TakeoffSummary) -> str:
        """Generates standard comma-separated measurement rows."""
        output = StringIO()
        writer = csv.writer(output)

        # Header metadata
        writer.writerow(["DRAWING TAKEOFF MEASUREMENT SHEET"])
        writer.writerow(["Drawing Number", takeoff.drawing_number, "Revision", takeoff.revision])
        writer.writerow([])

        # Table Header
        writer.writerow([
            "Item Code", "Description", "Location", "Net Quantity", "Unit",
            "Wastage %", "Gross Procurement Qty", "Formula Lineage", "Confidence", "Status"
        ])

        # Data Rows
        for item in takeoff.items:
            w_pct = getattr(item, "wastage_percent", 0.0) or 0.0
            g_qty = getattr(item, "gross_quantity", None)
            if g_qty is None:
                g_qty = round(item.quantity * (1.0 + w_pct), 2)

            unit_val = getattr(item.unit, "value", str(item.unit))
            status_val = getattr(item.status, "value", str(item.status))

            writer.writerow([
                item.item_code,
                item.description,
                item.location,
                round(item.quantity, 2),
                unit_val,
                f"{w_pct * 100:.1f}%",
                round(g_qty, 2),
                item.formula,
                f"{item.confidence * 100:.1f}%",
                status_val
            ])

        # Totals Section
        writer.writerow([])
        writer.writerow(["BOQ TOTALS SUMMARY"])
        writer.writerow(["Item Code", "Description", "Net Measured Qty", "Unit", "Wastage %", "Wastage Buffer", "Gross Procurement Qty", "Item Count"])
        totals = takeoff.total_by_item()
        for code, data in totals.items():
            net_q = data.get("net_quantity", data["total_quantity"])
            w_pct = data.get("wastage_percent", 0.0) or 0.0
            w_qty = data.get("wastage_quantity", 0.0)
            gross_q = data.get("gross_quantity", net_q)
            writer.writerow([
                code,
                data["description"],
                net_q,
                data["unit"],
                f"{w_pct * 100:.1f}%",
                w_qty,
                gross_q,
                data["count"]
            ])

        return output.getvalue()

    @staticmethod
    def export_excel(takeoff: TakeoffSummary, output_path: str | Path) -> Path:
        """Generates professionally styled Excel workbook with Takeoff and Summary sheets."""
        path = Path(output_path)
        path.parent.mkdir(parents=True, exist_ok=True)

        wb = openpyxl.Workbook()

        # Sheet 1: Measurement Takeoff
        ws1 = wb.active
        ws1.title = "Measurement Takeoff"

        # Sheet 2: BOQ Summary
        ws2 = wb.create_sheet(title="BOQ Summary")

        # Styling definitions
        font_title = Font(name="Segoe UI", size=14, bold=True, color="1F497D")
        font_header = Font(name="Segoe UI", size=10, bold=True, color="FFFFFF")
        fill_header = PatternFill(start_color="1F497D", end_color="1F497D", fill_type="solid")
        font_data = Font(name="Segoe UI", size=10)
        border_thin = Border(
            left=Side(style="thin", color="D9D9D9"),
            right=Side(style="thin", color="D9D9D9"),
            top=Side(style="thin", color="D9D9D9"),
            bottom=Side(style="thin", color="D9D9D9")
        )

        # Populate Sheet 1 (Takeoff)
        ws1["A1"] = "QUANTIFICATION ENGINE — DETAILED MEASUREMENT TAKEOFF"
        ws1["A1"].font = font_title
        ws1["A2"] = f"Drawing: {takeoff.drawing_number}  |  Revision: {takeoff.revision}  |  Standard: IS 1200 / POMI"
        ws1["A2"].font = Font(name="Segoe UI", size=10, italic=True, color="595959")

        headers_takeoff = [
            "Item Code", "Description", "Location", "Net Quantity", "Unit",
            "Wastage %", "Gross Qty", "Formula Lineage", "Confidence", "Status"
        ]

        for col_idx, h in enumerate(headers_takeoff, start=1):
            cell = ws1.cell(row=4, column=col_idx, value=h)
            cell.font = font_header
            cell.fill = fill_header
            cell.alignment = Alignment(horizontal="center" if col_idx in (1, 5, 6, 9, 10) else "left")

        current_row = 5
        for item in takeoff.items:
            w_pct = getattr(item, "wastage_percent", 0.0) or 0.0
            g_qty = getattr(item, "gross_quantity", None)
            if g_qty is None:
                g_qty = round(item.quantity * (1.0 + w_pct), 2)

            ws1.cell(row=current_row, column=1, value=item.item_code).alignment = Alignment(horizontal="center")
            ws1.cell(row=current_row, column=2, value=item.description)
            ws1.cell(row=current_row, column=3, value=item.location)
            ws1.cell(row=current_row, column=4, value=round(item.quantity, 2)).number_format = "#,##0.00"
            unit_val = getattr(item.unit, "value", str(item.unit))
            status_val = getattr(item.status, "value", str(item.status))
            ws1.cell(row=current_row, column=5, value=unit_val).alignment = Alignment(horizontal="center")
            ws1.cell(row=current_row, column=6, value=f"{w_pct*100:.1f}%").alignment = Alignment(horizontal="center")
            ws1.cell(row=current_row, column=7, value=round(g_qty, 2)).number_format = "#,##0.00"
            ws1.cell(row=current_row, column=8, value=item.formula)
            ws1.cell(row=current_row, column=9, value=f"{item.confidence*100:.1f}%").alignment = Alignment(horizontal="center")
            ws1.cell(row=current_row, column=10, value=status_val).alignment = Alignment(horizontal="center")

            for c in range(1, 11):
                cell = ws1.cell(row=current_row, column=c)
                cell.font = font_data
                cell.border = border_thin

            current_row += 1

        # Populate Sheet 2 (Summary)
        ws2["A1"] = "BOQ PROCUREMENT QUANTITY SUMMARY"
        ws2["A1"].font = font_title
        ws2["A2"] = f"Drawing: {takeoff.drawing_number}  |  Revision: {takeoff.revision}  |  Includes Trade Wastage Factors"
        ws2["A2"].font = Font(name="Segoe UI", size=10, italic=True, color="595959")

        headers_summary = [
            "Item Code", "Item Description", "Net Measured Qty", "Unit",
            "Wastage %", "Wastage Buffer", "Gross Procurement Qty", "Item Count"
        ]
        for col_idx, h in enumerate(headers_summary, start=1):
            cell = ws2.cell(row=4, column=col_idx, value=h)
            cell.font = font_header
            cell.fill = fill_header
            cell.alignment = Alignment(horizontal="center" if col_idx in (1, 4, 5, 8) else "left")

        current_row = 5
        totals = takeoff.total_by_item()
        for code, data in totals.items():
            net_q = data.get("net_quantity", data["total_quantity"])
            w_pct = data.get("wastage_percent", 0.0) or 0.0
            w_qty = data.get("wastage_quantity", 0.0)
            gross_q = data.get("gross_quantity", net_q)

            ws2.cell(row=current_row, column=1, value=code).alignment = Alignment(horizontal="center")
            ws2.cell(row=current_row, column=2, value=data["description"])
            ws2.cell(row=current_row, column=3, value=net_q).number_format = "#,##0.00"
            ws2.cell(row=current_row, column=4, value=data["unit"]).alignment = Alignment(horizontal="center")
            ws2.cell(row=current_row, column=5, value=f"{w_pct*100:.1f}%").alignment = Alignment(horizontal="center")
            ws2.cell(row=current_row, column=6, value=w_qty).number_format = "#,##0.00"
            ws2.cell(row=current_row, column=7, value=gross_q).number_format = "#,##0.00"
            ws2.cell(row=current_row, column=8, value=data["count"]).alignment = Alignment(horizontal="center")

            for c in range(1, 9):
                cell = ws2.cell(row=current_row, column=c)
                cell.font = font_data
                cell.border = border_thin
            current_row += 1

        # Sheet 3: Priced BOQ & Budget
        ws3 = wb.create_sheet(title="Priced BOQ & Budget")
        from pricing.cost_engine import CostEstimationEngine, FitoutGrade
        estimate = CostEstimationEngine.estimate_project_cost(takeoff, fitout_grade=FitoutGrade.STANDARD)

        ws3["A1"] = "PRICED BILL OF QUANTITIES & PROJECT COST ESTIMATE"
        ws3["A1"].font = font_title
        ws3["A2"] = f"Drawing: {takeoff.drawing_number}  |  Revision: {takeoff.revision}  |  Fitout Grade: Standard Commercial  |  Currency: INR (₹)"
        ws3["A2"].font = Font(name="Segoe UI", size=10, italic=True, color="595959")

        headers_priced = [
            "Item Code", "Item Description", "Trade Category", "Net Qty", "Unit",
            "Wastage %", "Gross PO Qty", "Material Rate (₹)", "Labor Rate (₹)",
            "Material Cost (₹)", "Labor Cost (₹)", "Total Item Cost (₹)"
        ]
        fill_priced_header = PatternFill(start_color="0B3C5D", end_color="0B3C5D", fill_type="solid")
        for col_idx, h in enumerate(headers_priced, start=1):
            cell = ws3.cell(row=4, column=col_idx, value=h)
            cell.font = font_header
            cell.fill = fill_priced_header
            cell.alignment = Alignment(horizontal="center" if col_idx in (1, 5, 6) else "left")

        start_item_row = 5
        cur_p_row = start_item_row
        for p_item in estimate.items:
            ws3.cell(row=cur_p_row, column=1, value=p_item.item_code).alignment = Alignment(horizontal="center")
            ws3.cell(row=cur_p_row, column=2, value=p_item.description)
            ws3.cell(row=cur_p_row, column=3, value=p_item.category)
            ws3.cell(row=cur_p_row, column=4, value=p_item.net_quantity).number_format = "#,##0.00"
            ws3.cell(row=cur_p_row, column=5, value=p_item.unit).alignment = Alignment(horizontal="center")
            ws3.cell(row=cur_p_row, column=6, value=f"{p_item.wastage_percent*100:.1f}%").alignment = Alignment(horizontal="center")
            ws3.cell(row=cur_p_row, column=7, value=p_item.gross_quantity).number_format = "#,##0.00"
            
            # Unit Rates
            base_mat_rate = p_item.material_cost_inr / max(0.001, p_item.gross_quantity)
            base_lab_rate = p_item.labor_cost_inr / max(0.001, p_item.net_quantity)
            ws3.cell(row=cur_p_row, column=8, value=round(base_mat_rate, 2)).number_format = "#,##0.00"
            ws3.cell(row=cur_p_row, column=9, value=round(base_lab_rate, 2)).number_format = "#,##0.00"

            # Formulas: Material Cost = Gross * MatRate, Labor Cost = Net * LabRate, Total = MatCost + LabCost
            ws3.cell(row=cur_p_row, column=10, value=f"=G{cur_p_row}*H{cur_p_row}").number_format = "₹#,##0.00"
            ws3.cell(row=cur_p_row, column=11, value=f"=D{cur_p_row}*I{cur_p_row}").number_format = "₹#,##0.00"
            ws3.cell(row=cur_p_row, column=12, value=f"=J{cur_p_row}+K{cur_p_row}").number_format = "₹#,##0.00"

            for c in range(1, 13):
                cell = ws3.cell(row=cur_p_row, column=c)
                cell.font = font_data
                cell.border = border_thin
            cur_p_row += 1

        end_item_row = max(start_item_row, cur_p_row - 1)

        # Financial Summary Section
        cur_p_row += 1
        font_subtotal = Font(name="Segoe UI", size=10, bold=True, color="1F497D")
        font_grand_total = Font(name="Segoe UI", size=11, bold=True, color="000000")
        fill_grand_total = PatternFill(start_color="D4EDDA", end_color="D4EDDA", fill_type="solid")

        # Direct Works Subtotal
        subtotal_row = cur_p_row
        ws3.cell(row=subtotal_row, column=11, value="Direct Works Subtotal:").font = font_subtotal
        ws3.cell(row=subtotal_row, column=11).alignment = Alignment(horizontal="right")
        ws3.cell(row=subtotal_row, column=12, value=f"=SUM(L{start_item_row}:L{end_item_row})").font = font_subtotal
        ws3.cell(row=subtotal_row, column=12).number_format = "₹#,##0.00"

        # Contractor OH&P (10%)
        ohp_row = subtotal_row + 1
        ws3.cell(row=ohp_row, column=11, value="Contractor OH & Profit (10%):").font = font_data
        ws3.cell(row=ohp_row, column=11).alignment = Alignment(horizontal="right")
        ws3.cell(row=ohp_row, column=12, value=f"=L{subtotal_row}*0.10").font = font_data
        ws3.cell(row=ohp_row, column=12).number_format = "₹#,##0.00"

        # Contingency Buffer (3%)
        cont_row = ohp_row + 1
        ws3.cell(row=cont_row, column=11, value="Contingency Buffer (3%):").font = font_data
        ws3.cell(row=cont_row, column=11).alignment = Alignment(horizontal="right")
        ws3.cell(row=cont_row, column=12, value=f"=L{subtotal_row}*0.03").font = font_data
        ws3.cell(row=cont_row, column=12).number_format = "₹#,##0.00"

        # Pre-Tax Total
        pre_tax_row = cont_row + 1
        ws3.cell(row=pre_tax_row, column=11, value="Subtotal (Before Tax):").font = font_subtotal
        ws3.cell(row=pre_tax_row, column=11).alignment = Alignment(horizontal="right")
        ws3.cell(row=pre_tax_row, column=12, value=f"=L{subtotal_row}+L{ohp_row}+L{cont_row}").font = font_subtotal
        ws3.cell(row=pre_tax_row, column=12).number_format = "₹#,##0.00"

        # GST @ 18%
        gst_row = pre_tax_row + 1
        ws3.cell(row=gst_row, column=11, value="GST @ 18%:").font = font_data
        ws3.cell(row=gst_row, column=11).alignment = Alignment(horizontal="right")
        ws3.cell(row=gst_row, column=12, value=f"=L{pre_tax_row}*0.18").font = font_data
        ws3.cell(row=gst_row, column=12).number_format = "₹#,##0.00"

        # GRAND TOTAL PROJECT BUDGET
        grand_row = gst_row + 1
        ws3.cell(row=grand_row, column=11, value="GRAND TOTAL ESTIMATE:").font = font_grand_total
        ws3.cell(row=grand_row, column=11).alignment = Alignment(horizontal="right")
        ws3.cell(row=grand_row, column=11).fill = fill_grand_total
        ws3.cell(row=grand_row, column=12, value=f"=L{pre_tax_row}+L{gst_row}").font = font_grand_total
        ws3.cell(row=grand_row, column=12).fill = fill_grand_total
        ws3.cell(row=grand_row, column=12).number_format = "₹#,##0.00"

        # Executive Commercial Fit-out Benchmark Metrics (Rate / Sqft & Rate / Sqm)
        total_sqm = float(takeoff.base_measurements.get("total_floor_area_sqm") or 0.0)
        if total_sqm <= 1.0:
            fl_sum = sum(it.quantity for it in takeoff.items if it.item_code.startswith("FL-") and getattr(it.unit, "value", str(it.unit)).lower() in ("sqm", "m2"))
            total_sqm = fl_sum if fl_sum > 0 else 1.0
        total_sqft = round(total_sqm * 10.7639104, 2)

        kpi_row1 = grand_row + 1
        ws3.cell(row=kpi_row1, column=11, value=f"Fit-out Rate / SQFT ({total_sqft:,.0f} sqft):").font = font_subtotal
        ws3.cell(row=kpi_row1, column=11).alignment = Alignment(horizontal="right")
        ws3.cell(row=kpi_row1, column=12, value=f"=L{grand_row}/{total_sqft:.2f}").font = font_subtotal
        ws3.cell(row=kpi_row1, column=12).number_format = "₹#,##0.00"

        kpi_row2 = kpi_row1 + 1
        ws3.cell(row=kpi_row2, column=11, value=f"Fit-out Rate / SQM ({total_sqm:,.1f} sqm):").font = font_subtotal
        ws3.cell(row=kpi_row2, column=11).alignment = Alignment(horizontal="right")
        ws3.cell(row=kpi_row2, column=12, value=f"=L{grand_row}/{total_sqm:.2f}").font = font_subtotal
        ws3.cell(row=kpi_row2, column=12).number_format = "₹#,##0.00"


        # Auto-adjust column widths for all sheets
        for sheet in (ws1, ws2, ws3):
            for col in sheet.columns:
                max_len = max(len(str(cell.value or "")) for cell in col)
                col_letter = get_column_letter(col[0].column)
                sheet.column_dimensions[col_letter].width = max(max_len + 3, 13)

        wb.save(str(path))
        return path
